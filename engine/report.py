import os
import sys
from datetime import datetime
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

# Ensure local imports work
sys.path.append(os.path.dirname(os.path.abspath(__file__)))
from compare import compare_pdfs

CURRENCY_SYMBOL = '$'

def create_report(results, output_path):
    wb = openpyxl.Workbook()
    
    # --- SHEET 1: Summary ---
    ws_sum = wb.active
    ws_sum.title = "Summary"
    
    # Calculate stats
    pos_checked = set()
    matched_count = 0
    total_risk = 0.0
    issues_per_po = {}
    
    for r in results:
        po = r['po_number']
        pos_checked.add(po)
        if po not in issues_per_po:
            issues_per_po[po] = 0
            
        if r['status'] == 'RECEIVED':
            matched_count += 1
        else:
            issues_per_po[po] += 1
            
        total_risk += r['amount_at_risk']
        
    total_rows = len(results)
    review_count = total_rows - matched_count
    
    ws_sum.append(["Warehouse Audit Summary"])
    ws_sum.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws_sum.append(["Generated On", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws_sum.append([])
    
    ws_sum.append(["Total POs Checked", len(pos_checked)])
    ws_sum.append(["Total Rows Matched", matched_count])
    ws_sum.append(["Total Rows Needing Review", review_count])
    ws_sum.append(["Total Amount at Risk", total_risk])
    ws_sum.cell(row=7, column=2).number_format = f'"{CURRENCY_SYMBOL}"#,##0.00'
    
    ws_sum.append([])
    ws_sum.append(["Breakdown by PO:"])
    ws_sum.cell(row=9, column=1).font = Font(bold=True)
    
    row_idx = 10
    for po in sorted(pos_checked):
        issues = issues_per_po.get(po, 0)
        if issues == 0:
            ws_sum.append([f"{po}: All items matched perfectly."])
        elif issues == 1:
            ws_sum.append([f"{po}: 1 item needs review."])
        else:
            ws_sum.append([f"{po}: {issues} items need review."])
        row_idx += 1
        
    ws_sum.column_dimensions['A'].width = 30
    ws_sum.column_dimensions['B'].width = 20

    # --- SHEET 2: Audit Details ---
    ws_det = wb.create_sheet("Audit Details")
    
    headers = ["PO Number", "Item Code", "Description", "Invoiced Qty", "Received Qty", "Unit Price", "Status", "Amount at Risk", "Action"]
    ws_det.append(headers)
    
    # Format Header
    for col_num, cell in enumerate(ws_det[1], 1):
        cell.font = Font(bold=True)
        
    ws_det.freeze_panes = 'A2'
    
    # Fills (Soft Colors)
    fill_green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    fill_amber = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    fill_red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    fill_grey = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    
    for r in results:
        action = "No action" if r['status'] == 'RECEIVED' else "Review"
        
        row = [
            r['po_number'],
            r['item_code'],
            r['description'],
            r['invoiced_qty'],
            r['received_qty'],
            r['unit_price'],
            r['status'],
            r['amount_at_risk'],
            action
        ]
        ws_det.append(row)
        
        curr_row_idx = ws_det.max_row
        
        # Format numbers/currency
        ws_det.cell(row=curr_row_idx, column=6).number_format = f'"{CURRENCY_SYMBOL}"#,##0.00'
        ws_det.cell(row=curr_row_idx, column=8).number_format = f'"{CURRENCY_SYMBOL}"#,##0.00'
        
        # Color row
        status = r['status']
        if status == 'RECEIVED':
            row_fill = fill_green
        elif status in ['SHORT_SHIPPED', 'OVER_SHIPPED']:
            row_fill = fill_amber
        elif status in ['MISSING', 'NOT_INVOICED']:
            row_fill = fill_red
        else:
            row_fill = fill_grey
            
        for col_num in range(1, len(headers) + 1):
            ws_det.cell(row=curr_row_idx, column=col_num).fill = row_fill

    # Auto-fit columns
    for col_idx, col in enumerate(ws_det.columns, 1):
        max_length = 0
        col_letter = get_column_letter(col_idx)
        for cell in col:
            try:
                if cell.value:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws_det.column_dimensions[col_letter].width = adjusted_width

    # --- SHEET 3: Supplier Messages ---
    ws_msg = wb.create_sheet("Supplier Messages")
    ws_msg.append(["PO Number", "Item Code", "Supplier Message"])
    for cell in ws_msg[1]:
        cell.font = Font(bold=True)
        
    for r in results:
        if r['status'] == 'RECEIVED':
            continue
            
        po = r['po_number']
        item = r['item_code']
        desc = r['description']
        inv = r['invoiced_qty']
        rec = r['received_qty']
        risk = f"{CURRENCY_SYMBOL}{r['amount_at_risk']:.2f}"
        
        if r['status'] == 'SHORT_SHIPPED':
            diff = inv - rec
            msg = f"Subject: {po} discrepancy, {desc}. Invoiced {inv}, received {rec} ({diff} short). Amount in question: {risk}. Please ship the balance or issue a credit note."
        elif r['status'] == 'OVER_SHIPPED':
            diff = rec - inv
            msg = f"Subject: {po} discrepancy, {desc}. Invoiced {inv}, received {rec} ({diff} over). Please advise if we should return the excess or if you will invoice us."
        elif r['status'] == 'MISSING':
            msg = f"Subject: {po} discrepancy, {desc}. Invoiced {inv}, but item was missing completely from shipment. Amount in question: {risk}. Please ship immediately or issue a credit note."
        elif r['status'] == 'NOT_INVOICED':
            msg = f"Subject: {po} discrepancy, {desc}. Item received (qty {rec}) but not on invoice. Please advise."
        elif r['status'] == 'PENDING':
            msg = f"Subject: {po} documentation missing. Invoice received but packing list is pending. Please provide the packing list to complete the audit."
        else:
            msg = "Subject: Discrepancy requires manual review."
            
        ws_msg.append([po, item, msg])
        
    ws_msg.column_dimensions['A'].width = 15
    ws_msg.column_dimensions['B'].width = 15
    ws_msg.column_dimensions['C'].width = 120
    
    # Save workbook
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path

if __name__ == '__main__':
    import shutil
    import json
    
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_dir = os.path.abspath(os.path.join(script_dir, '..'))
    input_dir = os.path.join(project_dir, 'input')
    processed_dir = os.path.join(project_dir, 'processed')
    report_dir = os.path.join(project_dir, 'reports')
    
    os.makedirs(processed_dir, exist_ok=True)
    
    results = compare_pdfs(input_dir)
    
    if not results:
        print(json.dumps({"error": "No data extracted. Ensure input folder has the required PDFs."}))
    else:
        output_path = os.path.join(report_dir, 'Audit_Report.xlsx')
        final_path = create_report(results, output_path)
        
        # Calculate summary stats for JSON response
        pos_checked = len(set(r['po_number'] for r in results))
        total_risk = sum(r['amount_at_risk'] for r in results)
        
        # Identify pending POs
        pending_pos = set(r['po_number'] for r in results if r['status'] == 'PENDING')
        
        # Move processed PDFs
        import glob
        for pdf_file in glob.glob(os.path.join(input_dir, '*.pdf')):
            filename = os.path.basename(pdf_file)
            # Check if this PDF belongs to a pending PO
            is_pending = any(po in filename for po in pending_pos)
            
            if not is_pending:
                shutil.move(pdf_file, os.path.join(processed_dir, filename))
                
        # Output strictly JSON for the Electron app to parse
        print(json.dumps({
            "posChecked": pos_checked,
            "amountAtRisk": total_risk,
            "reportPath": final_path
        }))
